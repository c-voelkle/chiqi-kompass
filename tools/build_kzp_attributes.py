"""Baut deploy/kzp24_attributes.js aus den BAG-Kennzahlen der Schweizer Spitäler.

Aufruf (Repo-Wurzel):
    python3 tools/build_kzp_attributes.py <kzp24_data.xlsx> <kzp24_KZ_TimeSerie.xlsx>

Quellen: BAG, Kennzahlen der Schweizer Spitäler 2024 (kzp24_data.xlsx, Blatt
KZ2024_KZP24) und Zeitreihe ab 2008 (kzp24_KZ_TimeSerie.xlsx). Übernommen
werden je Betrieb nur die Merkmale, nach denen das Cockpit Vergleichsgruppen
bildet: Kanton, Spitaltyp (BFS-Typologie), Rechtsform, Aktivitätstyp,
stationäre Austritte, Notfallstation/IPS, Aus- und Weiterbildung sowie
Akutbetten und CMI brutto (Akut). Betten und CMI sind in kzp24_data.xlsx
nicht befüllt; sie kommen aus der Zeitreihe, Datenjahr 2024, ersatzweise 2023
(das Jahr wird je Wert mitgeschrieben).

Zuordnung zu den Betrieben der QIP-Publikation (cockpit.html, #qip-data):
über den normalisierten Betriebsnamen. Namen, die sich zwischen den beiden
BAG-Publikationen unterscheiden, stehen ausdrücklich in ALIASES; jede
Zuordnung ist damit nachvollziehbar und wird nicht unscharf geraten.
Benötigt openpyxl.
"""
import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
COCKPIT = ROOT / 'deploy' / 'cockpit.html'
OUT = ROOT / 'deploy' / 'kzp24_attributes.js'

# QIP-Betriebsname -> Betriebsname in den Kennzahlen (geprüft anhand von Kanton und Typ).
ALIASES = {
    "Swiss Medical Network Hospitals SA Clinica Sant'Anna": "SMN - Clinica Sant'Anna",
    "Swiss Medical Network Hospitals SA Clinica Ars Medica": "SMN - Clinica Ars Medica",
    "Swiss Medical Network Hospitals SA Clinique Générale Ste-Anne": "SMN - Clinique Générale - Ste-Anne SA",
    "Swiss Medical Network Hospitals SA Clinique de Valère": "SMN - Clinique de Valère",
    "Swiss Medical Network Hospitals SA Privatklinik Obach": "SMN - Privatklinik Obach",
    "Swiss Medical Network Hospitals SA Neuchâtel": "SMN - Neuchâtel",
    "Swiss Medical Network Hospitals SA Privatklinik Siloah": "SMN - Privatklinik Siloah",
    "Swiss Medical Network Hospitals SA Privatklinik Villa im Park": "SMN - Klinik Villa im Park",
    "Swiss Medical Network Hospitals SA Privatklinik Belair": "SMN - Privatklinik Belair",
    "Pôle santé du Pays-d’Enhaut": "Pôle santé du Pays-d'Enhaut",
    "Genossenschaft Geburtshaus Simmental-Saanenland: Maternité Alpine":
        "Genossenschaft Geburtshaus Simmental-Saaneland - Maternité Alpine",
}


def norm(s):
    s = unicodedata.normalize('NFKD', s).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9]+', ' ', s).strip()


def tokens(value):
    return [t.strip() for t in str(value or '').split(',') if t.strip()]


def time_series(path):
    """{jahr: {normalisierter Name: Zeile}} aus dem Blatt «KZ ab … 2008»."""
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = next(w for w in wb.worksheets if w.title.startswith('KZ ab'))
    rows = ws.iter_rows(values_only=True)
    hdr = next(rows)
    out = {}
    for r in rows:
        rec = dict(zip(hdr, r))
        if rec.get('Inst') and rec.get('JAHR') in (2023, 2024):
            out.setdefault(rec['JAHR'], {})[norm(rec['Inst'])] = rec
    return out


def latest(ts, keys, column):
    """Wert aus 2024, ersatzweise 2023; (Wert, Jahr) oder (None, None)."""
    for year in (2024, 2023):
        for key in keys:
            rec = ts.get(year, {}).get(key)
            if rec and isinstance(rec.get(column), (int, float)):
                return rec[column], year
    return None, None


def main(xlsx, ts_path):
    html = COCKPIT.read_text(encoding='utf-8')
    data = json.loads(re.search(r'<script type="application/json" id="qip-data">(.*?)</script>', html, re.S).group(1))
    qip = [x['n'] for x in data['betrieb']['insts'] if not x['a']]

    ws = openpyxl.load_workbook(xlsx, read_only=True)['KZ2024_KZP24']
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    kzp = {}
    for r in rows[1:]:
        rec = dict(zip(hdr, r))
        if rec.get('Inst'):
            key = norm(rec['Inst'])
            if key in kzp:
                raise SystemExit('Doppelter Betriebsname in den Kennzahlen: ' + rec['Inst'])
            kzp[key] = rec

    ts = time_series(ts_path)
    out, missing = {}, []
    for name in qip:
        rec = kzp.get(norm(ALIASES.get(name, name)))
        if not rec:
            missing.append(name)
            continue
        aus = rec.get('AustritteT')
        out[name] = {
            'kt': rec['KT'],
            'typ': rec['Typ'] or None,
            'rf': rec['RForm'] or None,
            'akt': rec['Akt'] or None,
            'aus': int(aus) if isinstance(aus, (int, float)) else None,
            'sl': tokens(rec.get('SL')),
            'wb': tokens(rec.get('WB')),
            'kzp': rec['Inst'],
        }
        keys = [norm(rec['Inst']), norm(name)]
        bet, bet_year = latest(ts, keys, 'BettenStatA')
        cmi, cmi_year = latest(ts, keys, 'CMIbA')
        if bet is not None:
            out[name].update(bet=round(bet, 1), betJ=bet_year)
        if cmi is not None:
            out[name].update(cmi=round(cmi, 3), cmiJ=cmi_year)
    if missing:
        raise SystemExit('Ohne Zuordnung: ' + '; '.join(missing))

    payload = {
        'quelle': 'BAG, Kennzahlen der Schweizer Spitäler 2024 (kzp24_data.xlsx, Blatt KZ2024_KZP24; '
                  'Akutbetten und CMI aus kzp24_KZ_TimeSerie.xlsx, 2024, ersatzweise 2023)',
        'jahr': 2024,
        'betriebe': out,
    }
    OUT.write_text(
        '// Automatisch erzeugt von tools/build_kzp_attributes.py. Nicht von Hand editieren.\n'
        '// Quelle: ' + payload['quelle'] + '. Auszugsweiser Abdruck unter Quellenangabe gestattet.\n'
        'window.KZP24 = ' + json.dumps(payload, ensure_ascii=False, separators=(',', ':')) + ';\n',
        encoding='utf-8')
    with_bet = sum(1 for a in out.values() if 'bet' in a)
    with_cmi = sum(1 for a in out.values() if 'cmi' in a)
    print(f'{len(out)} von {len(qip)} QIP-Betrieben zugeordnet -> {OUT.relative_to(ROOT)}'
          f' (Akutbetten {with_bet}, CMI {with_cmi})')


if __name__ == '__main__':
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    main(sys.argv[1], sys.argv[2])
